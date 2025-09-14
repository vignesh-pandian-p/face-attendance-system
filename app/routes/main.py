from flask import Blueprint, render_template, request, redirect, url_for, send_file, make_response
import os
import cv2
import face_recognition
from datetime import datetime
from fpdf import FPDF
from app.services.face_recognition import get_known_faces, mark_attendance
from app.extensions import db
from app.models import User, AttendanceLog, Face, Class

main_bp = Blueprint('main', __name__)

UPLOAD_FOLDER = 'static/faces'
LOG_FOLDER = 'logs'  # Folder to store log files

@main_bp.route('/')
def index():
    return render_template('index.html')

import numpy as np
from app.models import Face

@main_bp.route('/classes/<int:class_id>/register', methods=['GET', 'POST'])
def register_student_in_class(class_id):
    class_obj = Class.query.get_or_404(class_id)
    if request.method == 'POST':
        name = request.form['name']
        if User.query.filter_by(name=name).first():
            return redirect(url_for('main.register_student_in_class', class_id=class_id))

        cap = cv2.VideoCapture(0)
        while True:
            ret, frame = cap.read()
            if not ret:
                break
            cv2.putText(frame, f"Registering for {class_obj.department} {class_obj.year}", (10, 30), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (0, 255, 0), 2)
            cv2.putText(frame, "Press 'S' to save, 'Q' to quit", (10, 60), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (0, 255, 0), 2)
            cv2.imshow("Capture Face", frame)
            key = cv2.waitKey(1) & 0xFF
            if key == ord('s'):
                new_user = User(name=name, class_id=class_id)
                db.session.add(new_user)
                db.session.commit()

                image_path = os.path.join(UPLOAD_FOLDER, f"{new_user.id}_{name}.jpg")
                cv2.imwrite(image_path, frame)

                rgb_frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
                encodings = face_recognition.face_encodings(rgb_frame)
                if encodings:
                    encoding_str = ','.join(map(str, encodings[0]))
                    new_face = Face(
                        user_id=new_user.id,
                        image_path=image_path,
                        encoding=encoding_str
                    )
                    db.session.add(new_face)
                    db.session.commit()
                break
            elif key == ord('q'):
                break
        cap.release()
        cv2.destroyAllWindows()
        return redirect(url_for('main.class_detail', class_id=class_id))
    return render_template('register.html', class_obj=class_obj)

from app.services.face_recognition import get_known_faces, mark_attendance

@main_bp.route('/attendance')
def attendance():
    known_encodings, known_names = get_known_faces()
    cap = cv2.VideoCapture(0)

    while True:
        ret, frame = cap.read()
        if not ret:
            break

        rgb_frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
        face_locations = face_recognition.face_locations(rgb_frame)
        face_encodings = face_recognition.face_encodings(rgb_frame, face_locations)

        for (top, right, bottom, left), face_encoding in zip(face_locations, face_encodings):
            matches = face_recognition.compare_faces(known_encodings, face_encoding)
            name = "Unknown"

            if True in matches:
                first_match_index = matches.index(True)
                name = known_names[first_match_index]
                mark_attendance(name)

            cv2.rectangle(frame, (left, top), (right, bottom), (0, 0, 255), 2)
            cv2.putText(frame, name, (left + 6, bottom - 6), cv2.FONT_HERSHEY_DUPLEX, 0.5, (255, 255, 255), 1)

        cv2.imshow('Attendance', frame)

        if cv2.waitKey(1) & 0xFF == ord('q'):
            break

    cap.release()
    cv2.destroyAllWindows()
    return redirect(url_for('main.index'))

from app.models import AttendanceLog

import io
import csv
from flask import Response

import openpyxl

@main_bp.route('/import', methods=['GET', 'POST'])
def import_excel():
    if request.method == 'POST':
        file = request.files['excel_file']
        if file:
            workbook = openpyxl.load_workbook(file)
            sheet = workbook.active
            for row in sheet.iter_rows(min_row=2, values_only=True):
                name, department, section, current_year = row
                if not User.query.filter_by(name=name).first():
                    # For now, image_path is a placeholder.
                    # The user will need to capture the face separately.
                    user = User(
                        name=name,
                        department=department,
                        section=section,
                        current_year=current_year,
                        image_path=f"static/faces/{name}.jpg" # Placeholder
                    )
                    db.session.add(user)
            db.session.commit()
            return redirect(url_for('main.dashboard'))
    return render_template('import_excel.html')

@main_bp.route('/add_face/<int:student_id>', methods=['GET'])
def add_face(student_id):
    student = User.query.get_or_404(student_id)
    cap = cv2.VideoCapture(0)
    while True:
        ret, frame = cap.read()
        if not ret:
            break
        cv2.putText(frame, f"Adding face for {student.name}", (10, 30), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (0, 255, 0), 2)
        cv2.putText(frame, "Press 'S' to save, 'Q' to quit", (10, 60), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (0, 255, 0), 2)
        cv2.imshow("Add Face", frame)
        key = cv2.waitKey(1) & 0xFF
        if key == ord('s'):
            image_path = os.path.join(UPLOAD_FOLDER, f"{student.id}_{student.name}_{len(student.faces) + 1}.jpg")
            cv2.imwrite(image_path, frame)

            rgb_frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
            encodings = face_recognition.face_encodings(rgb_frame)
            if encodings:
                encoding_str = ','.join(map(str, encodings[0]))
                new_face = Face(
                    user_id=student.id,
                    image_path=image_path,
                    encoding=encoding_str
                )
                db.session.add(new_face)
                db.session.commit()

            break
        elif key == ord('q'):
            break
    cap.release()
    cv2.destroyAllWindows()
    return redirect(url_for('main.edit_student', student_id=student_id))


@main_bp.route('/student/<int:student_id>/delete', methods=['POST'])
def delete_student(student_id):
    student = User.query.get_or_404(student_id)
    class_id = student.class_id

    # Delete associated face images
    for face in student.faces:
        if os.path.exists(face.image_path):
            os.remove(face.image_path)

    db.session.delete(student)
    db.session.commit()

    return redirect(url_for('main.class_detail', class_id=class_id))

@main_bp.route('/edit_student/<int:student_id>', methods=['GET', 'POST'])
def edit_student(student_id):
    student = User.query.get_or_404(student_id)
    if request.method == 'POST':
        student.name = request.form['name']
        student.class_id = request.form['class_id']
        db.session.commit()
        return redirect(url_for('main.class_detail', class_id=student.class_id))

    all_classes = Class.query.all()
    return render_template('edit_student.html', student=student, classes=all_classes)

@main_bp.route('/classes')
def classes():
    all_classes = Class.query.all()
    return render_template('classes.html', classes=all_classes)

@main_bp.route('/classes/add', methods=['GET', 'POST'])
def add_class():
    if request.method == 'POST':
        department = request.form['department']
        section = request.form['section']
        year = request.form['year']
        new_class = Class(department=department, section=section, year=year)
        db.session.add(new_class)
        db.session.commit()
        return redirect(url_for('main.classes'))
    return render_template('add_class.html')

@main_bp.route('/classes/<int:class_id>')
def class_detail(class_id):
    class_obj = Class.query.get_or_404(class_id)
    return render_template('class_detail.html', class_obj=class_obj)

@main_bp.route('/classes/<int:class_id>/attendance/download')
def download_class_attendance_pdf(class_id):
    class_obj = Class.query.get_or_404(class_id)

    student_ids = [user.id for user in class_obj.users]

    logs = AttendanceLog.query.filter(AttendanceLog.user_id.in_(student_ids)).order_by(AttendanceLog.timestamp.desc()).all()

    pdf = FPDF()
    pdf.add_page()

    pdf.set_font('Arial', 'B', 16)
    pdf.cell(0, 10, f'Attendance Report for {class_obj.year} {class_obj.department} - Section {class_obj.section}', 0, 1, 'C')
    pdf.ln(10)

    pdf.set_font('Arial', 'B', 12)
    pdf.cell(60, 10, 'Student Name', 1)
    pdf.cell(60, 10, 'Timestamp', 1)
    pdf.ln()

    pdf.set_font('Arial', '', 12)
    if logs:
        for log in logs:
            pdf.cell(60, 10, log.user.name, 1)
            pdf.cell(60, 10, log.timestamp.strftime('%Y-%m-%d %H:%M:%S'), 1)
            pdf.ln()
    else:
        pdf.cell(120, 10, 'No attendance records found for this class.', 1, 1, 'C')

    response = make_response(bytes(pdf.output(dest='S')))
    response.headers['Content-Type'] = 'application/pdf'
    response.headers['Content-Disposition'] = f'attachment; filename=attendance_report_{class_obj.department}_{class_obj.year}_{class_obj.section}.pdf'

    return response

@main_bp.route('/classes/<int:class_id>/delete', methods=['POST'])
def delete_class(class_id):
    class_to_delete = Class.query.get_or_404(class_id)
    if not class_to_delete.users:
        db.session.delete(class_to_delete)
        db.session.commit()
    return redirect(url_for('main.classes'))

@main_bp.route('/students')
def students():
    return redirect(url_for('main.classes'))

@main_bp.route('/dashboard')
def dashboard():
    logs = AttendanceLog.query.order_by(AttendanceLog.timestamp.desc()).all()
    return render_template('dashboard.html', logs=logs)

@main_bp.route('/download/attendance.csv')
def download_attendance_csv():
    logs = AttendanceLog.query.all()

    output = io.StringIO()
    writer = csv.writer(output)

    # Write headers
    writer.writerow(['ID', 'User Name', 'Timestamp'])

    # Write data
    for log in logs:
        writer.writerow([log.id, log.user.name, log.timestamp.strftime('%Y-%m-%d %H:%M:%S')])

    output.seek(0)

    return Response(
        output,
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment;filename=attendance.csv"}
    )
